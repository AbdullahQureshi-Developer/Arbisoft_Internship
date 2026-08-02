import io
import pytest
from datetime import datetime, timedelta
from docx import Document
from src.marshal.graph import process_slack_message
from src.skills.summarization_skill import SummaryResult
from src.skills.task_extraction_skill import TaskExtractionResult, ExtractedTask
from src.skills.date_extraction_skill import ExtractedDatesAndUpdates, ExtractedDateItem
from src.agents.calendar_agent import CreateEventResult


def test_document_upload_and_date_routing_flow(monkeypatch):
    now = datetime.utcnow()
    past_date = now - timedelta(hours=2)
    near_date = now + timedelta(hours=2) # <= 4 hours
    far_date = now + timedelta(days=2)   # > 4 hours

    # Create test DOCX document
    doc = Document()
    doc.add_heading("Sprint Planning Document", level=1)
    doc.add_paragraph(f"Past deadline: {past_date.isoformat()}")
    doc.add_paragraph(f"Near deadline: {near_date.isoformat()}")
    doc.add_paragraph(f"Far deadline: {far_date.isoformat()}")
    doc.add_paragraph("Notable update: Switched database to PostgreSQL.")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # Mocks
    monkeypatch.setattr("src.marshal.graph.summarize_text", lambda text: SummaryResult(summary="Sprint planning summary."))
    monkeypatch.setattr("src.marshal.graph.extract_tasks_and_reminders", lambda text: TaskExtractionResult(
        tasks=[ExtractedTask(description="Review spec", assignee="Alice")]
    ))
    
    mock_dates_and_updates = ExtractedDatesAndUpdates(
        dates=[
            ExtractedDateItem(description="Past Review", date=past_date),
            ExtractedDateItem(description="Near Reminder", date=near_date),
            ExtractedDateItem(description="Far Meeting", date=far_date),
        ],
        updates=["Switched database to PostgreSQL"],
    )
    monkeypatch.setattr("src.marshal.graph.extract_dates_and_updates", lambda text, reference_now=None: mock_dates_and_updates)

    scheduled_reminders = []
    def mock_schedule_reminder(req):
        scheduled_reminders.append(req)
        return type("MockRemRes", (), {"reminder_id": 1, "fire_at": req.message, "status": "scheduled"})()

    scheduled_calendar_events = []
    def mock_schedule_calendar(req):
        scheduled_calendar_events.append(req)
        return CreateEventResult(
            event_id="123",
            title=req.title,
            start_time=req.start_time,
            end_time=req.end_time,
            status="created",
            html_url="https://calendar.google.com/event?id=123",
            attendees=req.attendees,
        )

    monkeypatch.setattr("src.marshal.graph.schedule_new_reminder", mock_schedule_reminder)
    monkeypatch.setattr("src.marshal.graph.schedule_calendar_event", mock_schedule_calendar)

    response = process_slack_message(
        message_text="Here is the spec doc",
        channel_id="C12345",
        user_id="U_DOC_USER",
        file_bytes=docx_bytes,
        file_name="spec.docx",
    )

    # 1. Assert near date (<= 4h) created internal reminder
    assert len(scheduled_reminders) == 1
    assert scheduled_reminders[0].user_id == "U_DOC_USER"
    assert "Near Reminder" in scheduled_reminders[0].message
    assert scheduled_reminders[0].delay_seconds > 0

    # 2. Assert far date (> 4h) created Google Calendar event
    assert len(scheduled_calendar_events) == 1
    assert "Far Meeting" in scheduled_calendar_events[0].title

    # 3. Assert past date flagged as skipped in confirmation
    assert "already past — not scheduled" in response or "already past" in response
    assert "Past Review" in response

    # 4. Assert response details
    assert "Sprint planning summary" in response
    assert "Switched database to PostgreSQL" in response
    assert "Review spec" in response
